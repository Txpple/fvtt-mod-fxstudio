# The gap report as a workbook, grouped by compendium: one sheet per book, a Summary in front and
# an All sheet behind. Reads dist/gap-report.json (tools/gap-report.mjs writes it).
import json, os, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser('~'), 'Desktop', 'FX Studio gap analysis.xlsx')
rows = json.load(open(os.path.join(REPO, 'dist', 'gap-report.json'), encoding='utf-8'))

HEAD = ['Record', 'Type', 'On', 'What it does', 'Proposed FX', 'Based on', 'VFX (JB2A)', 'SFX (PSFX)', 'Confidence', 'Key']
WIDTH = [34, 9, 22, 40, 62, 20, 46, 44, 15, 34]

AMBER = PatternFill('solid', fgColor='F2C14E')
HEADFILL = PatternFill('solid', fgColor='2E2A26')
NONEFILL = PatternFill('solid', fgColor='F3F1EE')
CLOSE = PatternFill('solid', fgColor='E8F2E4')
FAMILY = PatternFill('solid', fgColor='FBF3DF')
THIN = Border(bottom=Side(style='thin', color='DDD8D0'))
WHITE = Font(color='FFFFFF', bold=True, size=11)
GREY = Font(color='8A8378')

def sheet_name(where):
    # a book's sheet name: short, unique, and legal (<=31 chars, no []:*?/\)
    n = where.replace('Forgotten Realms: ', '').replace('Ravenloft: The Horrors Within', 'Ravenloft')
    n = n.replace("Dungeon Master's Guide", 'DMG').replace("Player's Handbook", 'PHB').replace('Monster Manual', 'MM')
    n = n.replace(' · ', ' - ')
    for c in '[]:*?/\\':
        n = n.replace(c, '-')
    return n[:31]

def write_rows(ws, rs, show_where=False):
    head = (['Compendium'] if show_where else []) + HEAD
    width = ([30] if show_where else []) + WIDTH
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADFILL
        cell.font = WHITE
        cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = width[c - 1]
    ws.row_dimensions[1].height = 22
    for r in rs:
        vals = ([r['where']] if show_where else []) + [
            r['name'], r['kind'], r['on'], r['facts'], r['sentence'],
            r['basis'], r['vfx'], r['sfx'], r['confidence'], r['key'],
        ]
        ws.append(vals)
        i = ws.max_row
        fill = NONEFILL if r['confidence'] == 'No proposed FX' else (FAMILY if r['confidence'] == 'Family match' else CLOSE)
        for c in range(1, len(vals) + 1):
            cell = ws.cell(row=i, column=c)
            cell.border = THIN
            cell.alignment = Alignment(vertical='top', wrap_text=c in (4 + (1 if show_where else 0), 5 + (1 if show_where else 0)))
            if r['confidence'] == 'No proposed FX':
                cell.font = GREY
        ws.cell(row=i, column=len(vals) - 1).fill = fill
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(head))}{ws.max_row}'

wb = Workbook()

# ── Summary ──────────────────────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = 'Summary'
ws['A1'] = 'FX Studio — the gap: what the installed books hold that has no FX'
ws['A1'].font = Font(bold=True, size=15)
ws['A2'] = f'{len(rows)} addressed records with no FX. {sum(1 for r in rows if r["confidence"] != "No proposed FX")} have a proposed FX; {sum(1 for r in rows if r["confidence"] == "No proposed FX")} are marked no proposed FX.'
ws['A3'] = 'Every VFX and SFX path named here was checked against the installed JB2A Patreon 0.9.2 and PSFX databases — none is invented.'
ws['A4'] = '"Based on" names an FX the corpus already holds: copy it, then swap the asset shown. "Close match" = the record itself suggests it; "Family match" = it applies to any weapon, so the FX is a small mark rather than a swing.'
for r in (2, 3, 4):
    ws[f'A{r}'].font = GREY

start = 6
ws[f'A{start}'] = 'Compendium'
ws[f'B{start}'] = 'Gaps'
ws[f'C{start}'] = 'Proposed'
ws[f'D{start}'] = 'No proposed FX'
ws[f'E{start}'] = 'Sheet'
for c in 'ABCDE':
    ws[f'{c}{start}'].fill = HEADFILL
    ws[f'{c}{start}'].font = WHITE
by = collections.OrderedDict()
for r in rows:
    by.setdefault(r['where'], []).append(r)
order = sorted(by.items(), key=lambda kv: -len(kv[1]))
i = start + 1
for where, rs in order:
    n_yes = sum(1 for r in rs if r['confidence'] != 'No proposed FX')
    ws.cell(row=i, column=1, value=where)
    ws.cell(row=i, column=2, value=len(rs))
    ws.cell(row=i, column=3, value=n_yes)
    ws.cell(row=i, column=4, value=len(rs) - n_yes)
    ws.cell(row=i, column=5, value=sheet_name(where) if len(rs) >= 10 else ('This world' if where.startswith('this world') else 'Other books'))
    for c in range(1, 6):
        ws.cell(row=i, column=c).border = THIN
    i += 1
ws.cell(row=i, column=1, value='TOTAL').font = Font(bold=True)
ws.cell(row=i, column=2, value=len(rows)).font = Font(bold=True)
ws.cell(row=i, column=3, value=sum(1 for r in rows if r['confidence'] != 'No proposed FX')).font = Font(bold=True)
ws.cell(row=i, column=4, value=sum(1 for r in rows if r['confidence'] == 'No proposed FX')).font = Font(bold=True)
for c, w in zip('ABCDE', (48, 10, 12, 16, 26)):
    ws.column_dimensions[c].width = w
ws.freeze_panes = f'A{start + 1}'

# ── one sheet per book (>= 10 gaps), this world together, the rest in Other books ─────────────
world, other = [], []
for where, rs in order:
    rs = sorted(rs, key=lambda r: (r['kind'], r['name'].lower()))
    if where.startswith('this world'):
        world.extend(rs)
    elif len(rs) >= 10:
        write_rows(wb.create_sheet(sheet_name(where)), rs)
    else:
        other.extend(rs)
if world:
    write_rows(wb.create_sheet('This world'), sorted(world, key=lambda r: (r['where'], r['kind'], r['name'].lower())), show_where=True)
if other:
    write_rows(wb.create_sheet('Other books'), sorted(other, key=lambda r: (r['where'], r['kind'], r['name'].lower())), show_where=True)

# ── everything, in one sheet, for filtering ───────────────────────────────────────────────────
write_rows(wb.create_sheet('All gaps'), sorted(rows, key=lambda r: (r['where'], r['kind'], r['name'].lower())), show_where=True)

wb.save(OUT)
print(f'wrote {OUT}')
print(f'{len(wb.sheetnames)} sheets: {", ".join(wb.sheetnames)}')
